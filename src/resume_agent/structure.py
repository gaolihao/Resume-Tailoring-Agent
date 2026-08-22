from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path

from docx import Document
from docx.shared import Pt

from resume_agent.models import TailoredResume

SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "summary": ("professional summary", "summary", "profile", "objective"),
    "skills": ("skills", "technical skills", "core competencies", "technologies"),
    "experience": ("experience", "work experience", "employment", "work history"),
    "education": ("education", "academic background"),
    "additional": ("projects", "certifications", "additional", "awards", "publications"),
}


class SlotKind(str, Enum):
    HEADER = "header"
    SECTION_TITLE = "section_title"
    BODY = "body"
    ROLE = "role"
    BULLET = "bullet"


@dataclass
class ContentSlot:
    kind: SlotKind
    section: str | None
    original_text: str
    style: str | None = None
    bullet_prefix: str = "- "


@dataclass
class ResumeLayout:
    source_path: Path
    source_suffix: str
    slots: list[ContentSlot] = field(default_factory=list)
    docx_paragraphs: list | None = None


def load_layout(source_path: Path) -> ResumeLayout:
    suffix = source_path.suffix.lower()
    if suffix == ".docx":
        return _layout_from_docx(source_path)
    return _layout_from_text(source_path.read_text(encoding="utf-8"), source_path)


def _normalize_heading(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().lower().rstrip(":"))


def _detect_section(text: str) -> str | None:
    normalized = _normalize_heading(text)
    for section, aliases in SECTION_ALIASES.items():
        if normalized in aliases:
            return section
    return None


def _is_bullet(text: str) -> bool:
    return text.lstrip().startswith(("-", "•", "·", "*"))


def _bullet_prefix(text: str) -> str:
    stripped = text.lstrip()
    for prefix in ("- ", "• ", "· ", "* "):
        if stripped.startswith(prefix):
            return prefix
    if stripped.startswith("-") and not stripped.startswith("- "):
        return "- "
    return "- "


def _strip_bullet(text: str) -> str:
    return re.sub(r"^[\-\•\·\*]\s*", "", text.lstrip()).strip()


def _is_date_line(text: str) -> bool:
    if _is_bullet(text):
        return False
    has_year = bool(re.search(r"\b(19|20)\d{2}\b", text))
    has_title_marker = "—" in text or " - " in text
    return has_year and not has_title_marker and len(text) < 60


def _is_role_line(text: str) -> bool:
    if _is_bullet(text) or _is_date_line(text):
        return False
    lowered = text.lower()
    if "—" in text or " - " in text:
        return True
    if any(k in lowered for k in ("present", "–", "-", " to ")) and re.search(
        r"\b(19|20)\d{2}\b", text
    ):
        return True
    return False


def _layout_from_text(raw: str, source_path: Path) -> ResumeLayout:
    lines = raw.splitlines()
    slots: list[ContentSlot] = []
    current_section: str | None = None

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        section = _detect_section(stripped)
        if section:
            current_section = section
            slots.append(
                ContentSlot(
                    kind=SlotKind.SECTION_TITLE,
                    section=section,
                    original_text=stripped,
                )
            )
            continue

        if current_section is None:
            kind = SlotKind.HEADER
        elif current_section == "experience" and _is_bullet(stripped):
            kind = SlotKind.BULLET
        elif current_section == "experience" and _is_role_line(stripped):
            kind = SlotKind.ROLE
        elif current_section == "additional" and _is_bullet(stripped):
            kind = SlotKind.BULLET
        else:
            kind = SlotKind.BODY

        slots.append(
            ContentSlot(
                kind=kind,
                section=current_section,
                original_text=stripped,
                bullet_prefix=_bullet_prefix(stripped) if kind == SlotKind.BULLET else "- ",
            )
        )

    return ResumeLayout(
        source_path=source_path,
        source_suffix=source_path.suffix.lower(),
        slots=slots,
    )


def _layout_from_docx(source_path: Path) -> ResumeLayout:
    document = Document(str(source_path))
    slots: list[ContentSlot] = []
    current_section: str | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = paragraph.style.name if paragraph.style else None
        section = _detect_section(text)
        if section:
            current_section = section
            slots.append(
                ContentSlot(
                    kind=SlotKind.SECTION_TITLE,
                    section=section,
                    original_text=text,
                    style=style,
                )
            )
            continue

        is_list = style and "List" in style
        if current_section is None:
            kind = SlotKind.HEADER
        elif current_section == "experience" and (is_list or _is_bullet(text)):
            kind = SlotKind.BULLET
        elif current_section == "experience" and _is_role_line(text):
            kind = SlotKind.ROLE
        elif current_section == "additional" and (is_list or _is_bullet(text)):
            kind = SlotKind.BULLET
        else:
            kind = SlotKind.BODY

        slots.append(
            ContentSlot(
                kind=kind,
                section=current_section,
                original_text=text,
                style=style,
                bullet_prefix=_bullet_prefix(text) if kind == SlotKind.BULLET else "- ",
            )
        )

    return ResumeLayout(
        source_path=source_path,
        source_suffix=source_path.suffix.lower(),
        slots=slots,
        docx_paragraphs=list(document.paragraphs),
    )


def apply_tailored_text(slot: ContentSlot, new_text: str) -> str:
    if slot.kind == SlotKind.BULLET:
        return f"{slot.bullet_prefix}{new_text}"
    return new_text


def build_slot_texts(layout: ResumeLayout, tailored: TailoredResume) -> list[str | None]:
    texts: list[str | None] = [slot.original_text for slot in layout.slots]

    _fill_body_section(texts, layout, "summary", [tailored.professional_summary.strip()])
    _fill_body_section(texts, layout, "skills", [", ".join(tailored.skills_section)])
    _fill_experience_section(texts, layout, tailored.experience_entries)
    _fill_body_section(texts, layout, "education", tailored.education)
    _fill_bullet_section(texts, layout, "additional", tailored.additional_sections)

    return texts


def _indices_for(
    layout: ResumeLayout, section: str, kind: SlotKind
) -> list[int]:
    return [
        i
        for i, slot in enumerate(layout.slots)
        if slot.section == section and slot.kind == kind
    ]


def _fill_body_section(
    texts: list[str | None],
    layout: ResumeLayout,
    section: str,
    values: list[str],
) -> None:
    indices = _indices_for(layout, section, SlotKind.BODY)
    if not indices or not values:
        return

    texts[indices[0]] = values[0]
    for idx in indices[1:]:
        texts[idx] = ""


def _fill_bullet_section(
    texts: list[str | None],
    layout: ResumeLayout,
    section: str,
    values: list[str],
) -> None:
    indices = _indices_for(layout, section, SlotKind.BULLET)
    for idx, value in zip(indices, values, strict=False):
        texts[idx] = apply_tailored_text(layout.slots[idx], value)


def _experience_groups(layout: ResumeLayout) -> list[dict[str, list[int]]]:
    groups: list[dict[str, list[int]]] = []
    current: dict[str, list[int]] | None = None

    for i, slot in enumerate(layout.slots):
        if slot.section != "experience":
            continue
        if slot.kind == SlotKind.ROLE and _is_date_line(slot.original_text):
            if current is not None:
                current["roles"].append(i)
            continue
        if slot.kind == SlotKind.ROLE:
            current = {"roles": [i], "bullets": []}
            groups.append(current)
        elif slot.kind == SlotKind.BULLET and current is not None:
            current["bullets"].append(i)

    return groups


def _fill_experience_section(
    texts: list[str | None],
    layout: ResumeLayout,
    entries: list,
) -> None:
    groups = _experience_groups(layout)
    for group, entry in zip(groups, entries, strict=False):
        for idx, line in zip(group["roles"], entry.role_lines, strict=False):
            texts[idx] = line.strip()
        for idx, bullet in zip(group["bullets"], entry.bullets, strict=False):
            texts[idx] = apply_tailored_text(layout.slots[idx], bullet.strip())


def export_preserved_resume(
    source_path: Path,
    tailored: TailoredResume,
    output_path: Path,
) -> Path:
    layout = load_layout(source_path)
    texts = build_slot_texts(layout, tailored)
    suffix = layout.source_suffix

    if suffix == ".docx":
        return _write_docx_preserved(layout, texts, output_path)
    if suffix == ".pdf":
        return _write_pdf_preserved(layout, texts, output_path)
    return _write_text_preserved(source_path, layout, texts, output_path)


def _write_text_preserved(
    source_path: Path,
    layout: ResumeLayout,
    texts: list[str | None],
    output_path: Path,
) -> Path:
    original_lines = source_path.read_text(encoding="utf-8").splitlines()
    idx = 0
    rebuilt_lines: list[str] = []

    for line in original_lines:
        if not line.strip():
            rebuilt_lines.append(line)
            continue
        new_text = texts[idx] if idx < len(texts) else line.strip()
        idx += 1
        rebuilt_lines.append(new_text if new_text else "")

    if idx < len(texts):
        if rebuilt_lines and rebuilt_lines[-1].strip():
            rebuilt_lines.append("")
        for extra in texts[idx:]:
            if extra:
                rebuilt_lines.append(extra)

    output_path.write_text("\n".join(rebuilt_lines).strip() + "\n", encoding="utf-8")
    return output_path


def _rebuild_plain_text(layout: ResumeLayout, texts: list[str | None]) -> str:
    lines: list[str] = []
    for slot, text in zip(layout.slots, texts, strict=False):
        if text is None or text == "":
            continue
        if slot.kind == SlotKind.SECTION_TITLE and lines:
            lines.append("")
        lines.append(text)
    return "\n".join(lines).strip() + "\n"


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text


def _write_docx_preserved(
    layout: ResumeLayout,
    texts: list[str | None],
    output_path: Path,
) -> Path:
    shutil.copy2(layout.source_path, output_path)
    document = Document(str(output_path))
    non_empty = [p for p in document.paragraphs if p.text.strip()]

    if len(non_empty) != len(texts):
        return _write_docx_rebuilt(layout, texts, output_path)

    idx = 0
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        new_text = texts[idx]
        idx += 1
        if new_text is None:
            continue
        if new_text == "":
            _replace_paragraph_text(paragraph, "")
        else:
            _replace_paragraph_text(paragraph, new_text)

    document.save(str(output_path))
    return output_path


def _write_docx_rebuilt(
    layout: ResumeLayout,
    texts: list[str | None],
    output_path: Path,
) -> Path:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for slot, text in zip(layout.slots, texts, strict=False):
        if text is None or text == "":
            continue
        if slot.kind == SlotKind.SECTION_TITLE:
            document.add_paragraph(text)
        elif slot.kind == SlotKind.BULLET:
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    document.save(str(output_path))
    return output_path


def _write_pdf_preserved(
    layout: ResumeLayout,
    texts: list[str | None],
    output_path: Path,
) -> Path:
    try:
        import fitz  # pymupdf
    except ImportError as exc:
        raise RuntimeError(
            "PDF export requires pymupdf. Install with: pip install pymupdf"
        ) from exc

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    margin = 54
    y = margin
    line_height = 14
    page_height = 792 - margin

    for slot, text in zip(layout.slots, texts, strict=False):
        if text is None or text == "":
            continue

        fontsize = 11
        if slot.kind == SlotKind.SECTION_TITLE:
            fontsize = 12
            text = text.upper()
        elif slot.kind == SlotKind.HEADER and "@" in text:
            fontsize = 10

        if y + line_height > page_height:
            page = doc.new_page(width=612, height=792)
            y = margin

        page.insert_text((margin, y), text, fontsize=fontsize, fontname="helv")
        y += line_height + (2 if slot.kind == SlotKind.SECTION_TITLE else 0)

    doc.save(str(output_path))
    doc.close()
    return output_path
