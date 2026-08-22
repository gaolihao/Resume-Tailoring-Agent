from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

from resume_agent.models import JobAnalysis, TailoredResume
from resume_agent.structure import export_preserved_resume


def tailored_to_markdown(
    tailored: TailoredResume,
    job: JobAnalysis | None = None,
    *,
    include_audit: bool = True,
) -> str:
    lines: list[str] = []
    if job and include_audit:
        target = job.title
        if job.company:
            target = f"{job.title} @ {job.company}"
        lines.append(f"# Tailoring report — {target}")
        lines.append("")

    lines.extend(
        [
            "## Professional Summary",
            tailored.professional_summary,
            "",
            "## Skills",
            ", ".join(tailored.skills_section),
            "",
            "## Experience",
        ]
    )
    for entry in tailored.experience_entries:
        for role_line in entry.role_lines:
            lines.append(role_line)
        for bullet in entry.bullets:
            lines.append(f"- {bullet}")
        lines.append("")

    if tailored.education:
        lines.extend(["## Education"])
        for item in tailored.education:
            lines.append(f"- {item}")
        lines.append("")

    if tailored.additional_sections:
        lines.extend(["## Additional"])
        for item in tailored.additional_sections:
            lines.append(f"- {item}")
        lines.append("")

    if include_audit:
        lines.extend(["## Changes Made"])
        for change in tailored.changes_made:
            lines.append(f"- {change}")

        if tailored.keywords_injected:
            lines.extend(
                [
                    "",
                    "## Keywords Incorporated",
                    ", ".join(tailored.keywords_injected),
                ]
            )

    return "\n".join(lines).strip() + "\n"


def write_markdown(content: str, path: Path) -> Path:
    path.write_text(content, encoding="utf-8")
    return path


def export_tailored_resume(
    source_path: Path,
    tailored: TailoredResume,
    out_dir: Path,
    job: JobAnalysis | None = None,
) -> dict[str, Path]:
    """Write tailored resume in the original format plus an audit report."""
    out_dir.mkdir(parents=True, exist_ok=True)
    slug = _slug(job.title if job else "tailored")
    suffix = source_path.suffix.lower() or ".txt"

    primary_path = out_dir / f"resume_{slug}{suffix}"
    report_path = out_dir / f"resume_{slug}_report.md"

    export_preserved_resume(source_path, tailored, primary_path)
    write_markdown(
        tailored_to_markdown(tailored, job, include_audit=True),
        report_path,
    )
    return {"primary": primary_path, "report": report_path}


def write_docx(
    tailored: TailoredResume,
    path: Path,
    job: JobAnalysis | None = None,
) -> Path:
    """Legacy generic DOCX writer (prefer export_tailored_resume)."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if job:
        heading = doc.add_heading(job.title, level=1)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        if job.company:
            doc.add_paragraph(f"Target: {job.company}")

    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(tailored.professional_summary)

    doc.add_heading("Skills", level=2)
    doc.add_paragraph(", ".join(tailored.skills_section))

    doc.add_heading("Experience", level=2)
    for entry in tailored.experience_entries:
        for role_line in entry.role_lines:
            doc.add_paragraph(role_line)
        for bullet in entry.bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    if tailored.education:
        doc.add_heading("Education", level=2)
        for item in tailored.education:
            doc.add_paragraph(item, style="List Bullet")

    if tailored.additional_sections:
        doc.add_heading("Additional", level=2)
        for item in tailored.additional_sections:
            doc.add_paragraph(item, style="List Bullet")

    doc.save(str(path))
    return path


def _slug(value: str) -> str:
    cleaned = "".join(ch.lower() if ch.isalnum() else "_" for ch in value)
    while "__" in cleaned:
        cleaned = cleaned.replace("__", "_")
    return cleaned.strip("_")[:60] or "tailored"
